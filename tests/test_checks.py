"""Самопроверка модуля проверок на синтетических страницах (сайт не запрашивается).
Запуск: python -m unittest tests.test_checks -v"""
from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from techaudit.checks import checklist, report  # noqa: E402
from techaudit.checks.core import FAIL, NA, PASS, SKIPPED, WARN, CheckSettings, PageDoc, REGISTRY, role_of  # noqa: E402
from techaudit.checks.runner import load_all, load_results, run_checks  # noqa: E402
from techaudit.config import load_settings  # noqa: E402
from techaudit.storage import Store, now_iso  # noqa: E402

GOOD = """<!doctype html><html lang="ru"><head><meta charset="utf-8"><title>Товар</title>
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<link rel="canonical" href="https://shop.test/product/milk/">
<meta property="og:title" content="T"><meta property="og:description" content="D"><meta property="og:url" content="https://shop.test/product/milk/">
<meta property="og:image" content="https://shop.test/og.png"><meta property="og:type" content="website"><meta property="og:site_name" content="S">
<meta property="og:locale" content="ru_RU"><meta name="twitter:card" content="summary"><meta name="twitter:title" content="T">
<meta name="twitter:description" content="D"><meta name="twitter:image" content="https://shop.test/og.png">
<script type="application/ld+json">{"@context":"https://schema.org","@type":"Product","name":"Молоко","sku":123,
"offers":{"@type":"Offer","price":"59.9","priceCurrency":"RUB","availability":"https://schema.org/InStock"}}</script>
<script type="application/ld+json">{"@context":"https://schema.org","@type":"BreadcrumbList","itemListElement":[
{"@type":"ListItem","position":1,"name":"Главная"},{"@type":"ListItem","position":2,"name":"Молоко"}]}</script>
</head><body><header><nav><a href="/a">A</a><a href="/b">B</a><a href="/c">C</a></nav></header>
<div class="breadcrumbs"><a href="/">Главная</a><a href="/p">Молоко</a></div>
<h1>Молоко</h1><p>{text}</p><div>Цена 59,9 ₽</div>
<img src="/x.png" alt="молоко"><footer><a href="/f1">1</a><a href="/f2">2</a><a href="/f3">3</a></footer></body></html>"""
GOOD = GOOD.replace("{text}", "Описание товара. " * 60 + " ".join(f'<a href="/l{i}">l{i}</a>' for i in range(12)))

BAD = """<html><head><meta name="robots" content="noindex,nofollow"><meta name="viewport" content="width=device-width, maximum-scale=1">
<link rel="canonical" href="/Product/Milk"><link rel="canonical" href="https://other.test/x/"></head>
<body><span onclick="location.href='/x'">go</span><img src="/y.png"><a href="/p?PHPSESSID=abc">s</a></body></html>"""


class ChecksTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        load_all()
        cls.tmp = tempfile.TemporaryDirectory()
        cls.data = Path(cls.tmp.name) / "data"
        env = Path(cls.tmp.name) / ".env"
        env.write_text("NETWORKS=shop\nNET_SHOP_URL=https://shop.test\nNET_SHOP_TOKEN=t\n", encoding="utf-8")
        os.environ["DATA_DIR"] = str(cls.data)
        cls.settings = load_settings(str(env))
        cls.store = Store(cls.data)
        cls.store.start_run("r1", "prod", {})
        cls.add_page("product", "https://shop.test/product/milk/", GOOD, 120)
        cls.add_page("bad", "https://shop.test/Product/Milk", BAD, 900)
        cls.res = run_checks(cls.store, cls.settings, "r1", "shop")

    @classmethod
    def tearDownClass(cls):
        os.environ.pop("DATA_DIR", None)
        cls.store.close()
        cls.tmp.cleanup()

    @classmethod
    def add_page(cls, ptype, url, html, ttfb, variant="bot"):
        rel = cls.store.save_html("r1", "shop", ptype, url, variant, html.encode("utf-8"))
        cls.store.add_page(dict(run_id="r1", network="shop", page_type=ptype, variant=variant, ua_name="x", mode="raw",
                                source_url=url, url=url, status=200, final_url=url, chain="[]", headers="{}", ttfb_ms=ttfb,
                                total_ms=ttfb, bytes=len(html), sha256="x", file_path=rel, has_title=1, has_h1=1,
                                suspect="", error=None, fetched_at=now_iso()))

    def items(self, cid):
        return {i.kind: i for i in self.res[cid].items}

    def test_every_check_has_a_result_and_no_crash(self):
        self.assertEqual(set(self.res), set(REGISTRY))
        for cid, r in self.res.items():
            self.assertFalse(r.summary.startswith("ошибка проверки"), (cid, r.summary))

    def test_good_page_passes_and_bad_fails(self):
        for cid in ("canonical_present", "viewport", "og_url", "schema_product", "schema_breadcrumb", "lang", "charset"):
            it = self.items(cid)
            self.assertTrue(it["product"].ok, (cid, it["product"].msg))
        for cid in ("canonical_present", "viewport", "noindex_accidental", "nofollow_accidental", "lang", "charset",
                    "og_title_description", "links_a_href", "session_params"):
            self.assertFalse(self.items(cid)["bad"].ok, cid)

    def test_expects_noindex_page_is_not_flagged(self):
        self.assertEqual(role_of("search"), "search")

    def test_ttfb_uses_median_per_type(self):
        it = self.items("ttfb")
        self.assertTrue(it["product"].ok)
        self.assertFalse(it["bad"].ok)

    def test_checks_do_not_mutate_shared_document(self):
        # visible_text не должен удалять ld+json (раньше из-за этого пропадала разметка)
        p = PageDoc("shop", "product", "bot", "", "", "", 200, {}, 0, GOOD.encode(), "https://shop.test")
        _ = p.visible_text
        self.assertEqual(len(p.jsonld[0]), 2)

    def test_missing_probes_are_skipped_not_failed(self):
        for cid in ("redirect_slash", "status_404", "sitemap_available", "favicon"):
            self.assertEqual(self.res[cid].status, SKIPPED, cid)
        self.assertEqual(self.res["gsc"].status, SKIPPED)

    def test_probe_based_checks(self):
        st = self.store
        for kind, ver, reason in (("slash", "pass", ""), ("upper", "fail", "not_301:302")):
            st.add_probe(dict(run_id="r1", network="shop", kind=kind, page_type="product", source_url="x", url=f"https://shop.test/{kind}",
                              status=200, final_url="x", chain="[]", verdict=ver, reason=reason, fetched_at=now_iso()))
        st.add_probe(dict(run_id="r1", network="shop", kind="not_found", page_type="product", source_url="x",
                          url="https://shop.test/ta-404-1", status=200, final_url="https://shop.test/404", chain="[]",
                          verdict="fail", reason="redirected:302,status_200", fetched_at=now_iso()))
        res = run_checks(st, self.settings, "r1", "shop", only={"redirect_slash", "redirect_upper", "status_404", "status_no404_redirect"})
        self.assertEqual(res["redirect_slash"].status, PASS)
        self.assertEqual(res["redirect_upper"].status, FAIL)
        self.assertIn("not_301", res["redirect_upper"].summary)
        self.assertEqual(res["status_404"].status, FAIL)
        self.assertEqual(res["status_no404_redirect"].status, FAIL)

    def test_sitemap_completeness_and_lastmod(self):
        st = self.store
        xml = ("<?xml version='1.0'?><urlset xmlns='http://www.sitemaps.org/schemas/sitemap/0.9'>"
               "<url><loc>https://shop.test/product/milk/</loc><lastmod>2026-10-01</lastmod></url></urlset>")
        rel = st.save_blob("r1", "shop", "_probes/sitemap", "https://shop.test/sm.xml", "xml", xml.encode())
        st.add_probe(dict(run_id="r1", network="shop", kind="sitemap", url="https://shop.test/sm.xml", status=200, final_url="x",
                          chain="[]", file_path=rel, verdict="pass", reason="", meta=json.dumps({"type": "urlset", "urls": 1, "bytes": 200}),
                          fetched_at=now_iso()))
        res = run_checks(st, self.settings, "r1", "shop", only={"sitemap_complete", "sitemap_valid", "sitemap_lastmod"})
        self.assertEqual(res["sitemap_valid"].status, PASS)
        self.assertEqual(res["sitemap_complete"].status, FAIL)  # /Product/Milk (bad) нет в карте
        self.assertEqual(res["sitemap_lastmod"].status, WARN)  # одна дата у всех

    def test_xlsx_and_html_reports(self):
        import openpyxl
        src = Path(self.tmp.name) / "list.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Блок", "Пункт проверки", "Критерий", "Приоритет", "Проверяющий", "Комментарий"])
        ws.append(["Canonical", "Присутствует на всех индексируемых страницах", "один тег", "Критично", "a@b", "мой коммент"])
        ws.append(["Прочее", "Нечто ручное", "-", "Средний", "a@b", None])
        wb.save(src)
        run_checks(self.store, self.settings, "r1", "shop")  # полный пересчёт: прочие тесты считали подмножества
        results = load_results(self.store, "r1")
        dst = Path(self.tmp.name) / "out.xlsx"
        info = checklist.write_xlsx(src, dst, results, "r1")
        self.assertEqual(info["unmatched"], ["Нечто ручное"])
        out = openpyxl.load_workbook(dst)
        ws2 = out.worksheets[0]
        self.assertEqual(ws2.cell(2, 6).value, "мой коммент")  # исходные ячейки не тронуты
        self.assertEqual(ws2.cell(2, 7).value, "Не пройдено")
        self.assertIn("Автопроверка - детали", out.sheetnames)
        rows = checklist.read_rows(src)
        html = report.build_html("r1", rows, results, {})
        self.assertIn("Присутствует на всех индексируемых страницах", html)


class FakeOpenAI:
    """Мини-сервер, имитирующий /v1/chat/completions: запоминает запросы, отвечает JSON по типу промпта."""

    def __init__(self):
        import threading
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        outer = self
        self.requests: list[dict] = []

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def do_POST(self):  # noqa: N802
                body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                outer.requests.append({"auth": self.headers.get("Authorization"), "path": self.path, "body": body})
                user = body["messages"][1]["content"]
                if "результаты автопроверок" in user:
                    ids = [c["id"] for c in json.loads(user.split("ДАННЫЕ:\n", 1)[1])]
                    content = json.dumps({"summary": "Сводка теста.", "root_causes": [{"title": "Причина", "checks": ids[:1], "action": "Исправить"}],
                                          "tasks": [{"title": "Задача", "priority": "Критично", "checks": ids[:1], "description": "d"}],
                                          "comments": {i: f"комментарий {i}" for i in ids}}, ensure_ascii=False)
                    content = "<think>рассуждаю</think>```json\n" + content + "\n```"
                else:
                    pages = json.loads(user.split("ДАННЫЕ:\n", 1)[1])
                    content = json.dumps({"pages": [{"url": p["url"], "has_description": False, "hidden": False, "comment": "нет"} for p in pages]})
                out = json.dumps({"choices": [{"message": {"content": content}}]}).encode()
                self.send_response(200)
                self.send_header("Content-Length", str(len(out)))
                self.end_headers()
                self.wfile.write(out)

        self.srv = ThreadingHTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()
        self.url = f"http://127.0.0.1:{self.srv.server_address[1]}/v1"


class AIStageTest(ChecksTest):
    def test_ai_stage_adds_notes_without_changing_statuses(self):
        from techaudit.checks.ai_stage import load_notes, run_ai
        from techaudit.checks.ctx import Ctx
        import openpyxl
        run_checks(self.store, self.settings, "r1", "shop")
        before = {c: r["status"] for c, r in load_results(self.store, "r1")["shop"].items()}
        fake = FakeOpenAI()
        saved = dict(os.environ)
        os.environ.update({"AI_API_KEY": "k-test", "AI_API_URL": fake.url, "AI_MODEL": "x5-airun-medium"})
        try:
            settings = load_settings(None)  # ключ из окружения процесса
            ctx = Ctx(self.store, settings, "r1", "shop", CheckSettings.from_env(), "https://shop.test")
            results = load_results(self.store, "r1")["shop"]
            rows = [{"row": 2, "block": "B", "title": "t", "criterion": "", "priority": "Критично", "check_id": c} for c in results]
            st = run_ai(self.store, settings, ctx, results, rows, "r1", "shop", quiet=True)
            self.assertEqual(st["status"], "ok", st)
            n_first = len(fake.requests)
            self.assertEqual(n_first, 2)  # анализ + описания
            req = fake.requests[0]
            self.assertEqual(req["auth"], "Bearer k-test")
            self.assertTrue(req["path"].endswith("/chat/completions"))
            self.assertEqual(req["body"]["model"], "x5-airun-medium")
            self.assertNotIn("X-SP-Token", json.dumps(req["body"]))  # токен сайта не уходит в модель
            # повторный запуск — из кеша, без запросов к API
            run_ai(self.store, settings, ctx, results, rows, "r1", "shop", quiet=True)
            self.assertEqual(len(fake.requests), n_first)
            # статусы не изменились, выводы сохранены и попали в отчёты
            after = {c: r["status"] for c, r in load_results(self.store, "r1")["shop"].items()}
            self.assertEqual(before, after)
            notes = load_notes(self.store, "r1")
            self.assertEqual(notes["shop"]["analysis"]["summary"], "Сводка теста.")
            html = report.build_html("r1", rows, load_results(self.store, "r1"), {}, notes)
            self.assertIn("Сводка ИИ", html)
            src = Path(self.tmp.name) / "l2.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["Блок", "Пункт", "Критерий", "Приоритет", "П", "К"])
            wb.active.append(["Canonical", "Присутствует на всех индексируемых страницах", "", "Критично", "", ""])
            wb.save(src)
            dst = Path(self.tmp.name) / "o2.xlsx"
            checklist.write_xlsx(src, dst, load_results(self.store, "r1"), "r1", notes)
            out = openpyxl.load_workbook(dst)
            self.assertIn("ИИ - сводка", out.sheetnames)
            self.assertIn("комментарий canonical_present", str(out.worksheets[0].cell(2, 12).value))
            fake.srv.shutdown()
        finally:
            os.environ.clear()
            os.environ.update(saved)

    def test_no_api_key_is_skipped_not_error(self):
        from techaudit.checks.ai_stage import run_ai
        from techaudit.checks.ctx import Ctx
        saved = dict(os.environ)
        os.environ.pop("AI_API_KEY", None)
        try:
            settings = load_settings(None)
            ctx = Ctx(self.store, settings, "r1", "shop", CheckSettings.from_env(), "https://shop.test")
            st = run_ai(self.store, settings, ctx, {}, [], "r1", "shop", quiet=True)
            self.assertEqual(st["status"], "skipped")
            self.assertIn("AI_API_KEY", st["reason"])
        finally:
            os.environ.clear()
            os.environ.update(saved)

    def test_request_limit_stops_ai_without_crash(self):
        from techaudit.checks.ai_stage import run_ai
        from techaudit.checks.ctx import Ctx
        fake = FakeOpenAI()
        saved = dict(os.environ)
        os.environ.update({"AI_API_KEY": "k", "AI_API_URL": fake.url, "AI_MAX_REQUESTS": "0", "AI_MODEL": "limit-model"})
        try:
            settings = load_settings(None)
            ctx = Ctx(self.store, settings, "r1", "shop", CheckSettings.from_env(), "https://shop.test")
            res = {"viewport": {"status": "fail", "summary": "s", "items": [], "source": "html"}}
            st = run_ai(self.store, settings, ctx, res, [], "r1", "shop", quiet=True)
            self.assertEqual(st["status"], "error")
            self.assertIn("AI_MAX_REQUESTS", st["reason"])
            self.assertEqual(fake.requests, [])
            fake.srv.shutdown()
        finally:
            os.environ.clear()
            os.environ.update(saved)


class EndToEndMockTest(unittest.TestCase):
    """collect_html.py → run_checks.py --live на мок-сайте: пробы, живая выборка, отчёты."""

    def test_collect_then_checks_with_live(self):
        import io
        from contextlib import redirect_stdout
        from tests import mock_site
        from techaudit.collect import main as collect_main
        from techaudit.checks.cli import main as checks_main
        srv, base = mock_site.start()
        tmp = tempfile.TemporaryDirectory()
        t = Path(tmp.name)
        (t / "registry").mkdir()
        (t / "registry" / "mock.csv").write_text("type,url,note\nmain,/,\ncatalog,/catalog/,\nssr,/ssr/,\nsub,/catalog/sub/,\n", encoding="utf-8")
        saved = dict(os.environ)
        os.environ.update({"NETWORKS": "mock", "NET_MOCK_URL": base, "NET_MOCK_TOKEN": mock_site.TOKEN,
                           "DATA_DIR": str(t / "data"), "REGISTRY_DIR": str(t / "registry"), "UA_PROFILE": "googlebot"})
        try:
            with redirect_stdout(io.StringIO()):
                self.assertEqual(collect_main(["--variants", "plain,bot", "--probes", "all", "--no-diff"]), 0)
                code = checks_main(["--live", "--checklist", str(t / "none.xlsx")])
            self.assertEqual(code, 0)
            st = Store(t / "data")
            res = load_results(st, st.db.execute("SELECT run_id FROM runs").fetchone()[0])["mock"]
            slash = {i["kind"]: i for i in res["redirect_slash"]["items"]}
            self.assertTrue(any(i["ok"] and "/catalog" in i["url"] for i in res["redirect_slash"]["items"]), slash)
            self.assertEqual(res["redirect_upper"]["status"], FAIL)  # мок отдаёт 302
            self.assertEqual(res["status_no404_redirect"]["status"], FAIL)
            self.assertIn(res["links_404"]["status"], (PASS, FAIL))  # живая выборка выполнена, а не «пропущено»
            self.assertNotEqual(res["links_redirects"]["status"], SKIPPED)
            self.assertNotEqual(res["sitemap_available"]["status"], SKIPPED)
            self.assertEqual(res["sitemap_valid"]["status"], PASS)
            self.assertTrue(list((t / "data" / "reports").glob("*-checks.html")))
            st.close()
        finally:
            srv.shutdown(); srv.server_close()
            os.environ.clear(); os.environ.update(saved)
            tmp.cleanup()


class NoSecretsTest(unittest.TestCase):
    def test_checks_package_does_not_read_secret_env_directly(self):
        # токен и UA — только через config.py (правило проекта: данные в env, один источник)
        src = "".join(p.read_text(encoding="utf-8") for p in (Path(__file__).resolve().parent.parent / "techaudit" / "checks").glob("*.py"))
        for bad in ("SP_TOKEN", "NET_", "_TOKEN"):
            self.assertNotIn(f"environ.get(\"{bad}", src)


if __name__ == "__main__":
    unittest.main()
