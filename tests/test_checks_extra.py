"""Тесты дополнительных проверок и плана работ (приоритет, роль, часы). Запуск: python -m unittest tests.test_checks_extra -v"""
from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from techaudit.checks import checklist, plan, report  # noqa: E402
from techaudit.checks.core import FAIL, NA, PASS, SKIPPED, WARN, REGISTRY  # noqa: E402
from techaudit.checks.extras import EXTRA, with_extras  # noqa: E402
from techaudit.checks.live import tls_info  # noqa: E402
from techaudit.checks.runner import load_all, load_results, run_checks  # noqa: E402
from techaudit.checks.x_robots import Robots  # noqa: E402
from techaudit.checks.x_server import parse_cache_control  # noqa: E402
from techaudit.config import load_settings  # noqa: E402
from techaudit.storage import Store, now_iso  # noqa: E402

BIG = " ".join(f"слово{i}" for i in range(300))
LONG = "Это длинное описание страницы с достаточным количеством слов для проверки совпадения основного текста между страницами. " * 12


def page(title, desc, h1, body="", head="", links=""):
    d = f'<meta name="description" content="{desc}">' if desc is not None else ""
    h = "".join(f"<h1>{x}</h1>" for x in h1)
    return (f'<!doctype html><html lang="ru"><head><meta charset="utf-8"><title>{title}</title>{d}{head}</head>'
            f"<body><header><nav><a href='/'>Главная</a></nav></header><main>{h}<p>{body}</p>{links}</main><footer>© shop</footer></body></html>")


class ExtraChecksTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        load_all()
        cls.tmp = tempfile.TemporaryDirectory()
        cls.data = Path(cls.tmp.name) / "data"
        env = Path(cls.tmp.name) / ".env"
        env.write_text("NETWORKS=shop\nNET_SHOP_URL=https://shop.test\nNET_SHOP_TOKEN=t\n", encoding="utf-8")
        os.environ["DATA_DIR"] = str(cls.data)
        cls.settings = load_settings(str(env))
        cls.store = Store(cls.data)
        cls.store.start_run("r1", "prod", {})
        T = "Молоко 3,2% купить в магазине с доставкой"
        D = "Молоко коровье пастеризованное, доставка на дом и самовывоз, актуальные цены и акции на сегодня."
        add = cls.add_page
        add("a", "https://shop.test/milk/", page(T, D, ["Молоко"], BIG, '<link rel="canonical" href="https://shop.test/milk/">'),
            {"content-encoding": "gzip", "cache-control": "public, max-age=3600", "strict-transport-security": "max-age=31536000",
             "x-content-type-options": "nosniff", "referrer-policy": "no-referrer", "x-frame-options": "DENY", "etag": "\"1\""})
        # дубль title, заглушка в description, uppercase и utm в ссылках, смешанный контент
        add("b", "https://shop.test/bread/", page(T, "metaDescription", ["Хлеб"], BIG + " новый хвост про хлеб",
            '<link rel="canonical" href="https://shop.test/bread/"><script src="http://cdn.test/x.js"></script>',
            links="<a href='/Catalog/Milk/?utm_source=x'>подробнее</a><a href='http://shop.test/a/'>здесь</a><a href='/x//y/'></a>"),
            {"cache-control": "public, max-age max-age=36000", "vary": "Accept-Encoding"})
        # два H1, длинный title, нет description; canonical-цепочка d -> e -> f
        add("c", "https://shop.test/cheese/", page("С" * 95, None, ["Сыр", "Сыр ещё"], "коротко",
            '<link rel="canonical" href="https://shop.test/cheese-new/">'), {})
        add("e", "https://shop.test/cheese-new/", page("Сыр новый и достаточно длинный title", "Описание сыра достаточной длины для проверки, ещё немного слов тут.",
            ["Сыр новый"], "коротко", '<link rel="canonical" href="https://shop.test/cheese-final/">'), {})
        add("t", "https://shop.test/crm-test/", page("Тест страница для проверки", "Описание тестовой страницы для проверки, достаточно длинное.",
            ["Тест"], "x", '<link rel="canonical" href="https://shop.test/crm-test/">'), {})
        add("s", "https://shop.test/missing/", page("Страница не найдена", "Описание страницы с ошибкой, достаточно длинное для проверки.",
            ["404"], "x", '<link rel="canonical" href="https://shop.test/missing/">'), {})
        # plain-вариант с другим title и без Vary
        cls.add_page("a", "https://shop.test/milk/", page("Другой title для обычного клиента", "d", ["Молоко"], LONG), {}, variant="plain")
        cls.add_probe_file("robots", "https://shop.test/robots.txt", "txt",
                           "User-agent: *\nDisallow: /Catalog/\nDisallow: /crm-test/$\nUser-agent: GPTBot\nDisallow: /\nHost: shop.test\n", "text/plain")
        xml = ("<?xml version='1.0'?><urlset xmlns='http://www.sitemaps.org/schemas/sitemap/0.9'>"
               "<url><loc>https://shop.test/milk/</loc></url><url><loc>https://shop.test/Catalog/undefined/</loc></url>"
               "<url><loc>http://other.test/x/</loc></url><url><loc>https://shop.test/crm-test/</loc></url>"
               "<url><loc>https://shop.test/milk/</loc></url></urlset>")
        rel = cls.store.save_blob("r1", "shop", "_probes/sitemap", "https://shop.test/sm.xml", "xml", xml.encode())
        cls.store.add_probe(dict(run_id="r1", network="shop", kind="sitemap", url="https://shop.test/sm.xml", status=200, final_url="x",
                                 chain="[]", file_path=rel, verdict="pass", reason="",
                                 meta=json.dumps({"type": "urlset", "urls": 5, "bytes": 500}), fetched_at=now_iso()))
        cls.res = run_checks(cls.store, cls.settings, "r1", "shop")

    @classmethod
    def tearDownClass(cls):
        os.environ.pop("DATA_DIR", None)
        cls.store.close()
        cls.tmp.cleanup()

    @classmethod
    def add_page(cls, ptype, url, html, headers, variant="bot"):
        rel = cls.store.save_html("r1", "shop", ptype, url, variant, html.encode("utf-8"))
        cls.store.add_page(dict(run_id="r1", network="shop", page_type=ptype, variant=variant, ua_name="x", mode="raw",
                                source_url=url, url=url, status=200, final_url=url, chain="[]", headers=json.dumps(headers),
                                ttfb_ms=100, total_ms=100, bytes=len(html), sha256="x", file_path=rel, has_title=1, has_h1=1,
                                suspect="", error=None, fetched_at=now_iso()))

    @classmethod
    def add_probe_file(cls, kind, url, ext, text, ctype):
        rel = cls.store.save_blob("r1", "shop", f"_probes/{kind}", url, ext, text.encode())
        cls.store.add_probe(dict(run_id="r1", network="shop", kind=kind, url=url, status=200, final_url=url, chain="[]", file_path=rel,
                                 content_type=ctype, verdict="pass", reason="", fetched_at=now_iso()))

    def urls(self, cid, ok=None):
        return {i.url for i in self.res[cid].items if ok is None or i.ok == ok}

    # ---- реестр и план
    def test_all_checks_ran_without_crash_and_have_plan(self):
        self.assertEqual(set(self.res), set(REGISTRY))
        for cid, r in self.res.items():
            self.assertFalse(r.summary.startswith("ошибка проверки"), (cid, r.summary))
            self.assertIn(cid, plan.PLAN, cid)
        self.assertGreaterEqual(len(EXTRA), 30)

    # ---- мета и дубли
    def test_duplicate_title_found(self):
        self.assertEqual(self.res["dup_title"].status, FAIL)
        self.assertEqual(self.urls("dup_title", ok=False), {"https://shop.test/milk/", "https://shop.test/bread/"})

    def test_placeholders_title_description_h1(self):
        self.assertEqual(self.res["meta_placeholders"].status, FAIL)
        self.assertIn("https://shop.test/bread/", self.urls("meta_placeholders", ok=False))
        self.assertNotIn("https://shop.test/milk/", self.urls("meta_placeholders", ok=False))

    def test_title_description_h1_quality(self):
        self.assertIn("https://shop.test/cheese/", self.urls("title_quality", ok=False))
        self.assertEqual(self.res["description_quality"].status, FAIL)  # у сыра нет description
        self.assertIn("https://shop.test/cheese/", self.urls("description_quality", ok=False))
        h1 = {i.url: i for i in self.res["h1_single"].items}
        self.assertFalse(h1["https://shop.test/cheese/"].ok)
        self.assertTrue(h1["https://shop.test/milk/"].ok)

    def test_canonical_chain_and_soft_404(self):
        self.assertEqual(self.res["canonical_chain"].status, FAIL)
        self.assertIn("https://shop.test/cheese/", self.urls("canonical_chain", ok=False))
        self.assertIn("https://shop.test/missing/", self.urls("soft_404", ok=False))

    def test_duplicate_content_near_copy(self):
        r = self.res["duplicate_content"]
        self.assertEqual(r.status, FAIL)
        self.assertIn("https://shop.test/bread/", self.urls("duplicate_content", ok=False))

    # ---- robots.txt
    def test_robots_parser_semantics(self):
        rb = Robots("User-agent: *\nDisallow: /a\nAllow: /a/b\nDisallow: /*.json$\nDisallow:\nClean-param: utm_source\n"
                    "User-agent: YandexBot\nDisallow: /y/\nSitemap: https://x.test/sm.xml\n")
        self.assertEqual(rb.check("Googlebot", "https://x.test/a/c")[0], False)
        self.assertEqual(rb.check("Googlebot", "https://x.test/a/b/c")[0], True)  # Allow длиннее Disallow
        self.assertEqual(rb.check("Googlebot", "https://x.test/d/file.json")[0], False)
        self.assertEqual(rb.check("Googlebot", "https://x.test/d/file.json?x=1")[0], True)  # $ — конец адреса
        self.assertEqual(rb.check("YandexBot", "https://x.test/a")[0], True)  # своя группа вытесняет *
        self.assertEqual(rb.check("YandexBot", "https://x.test/y/z")[0], False)
        self.assertEqual(rb.sitemaps, ["https://x.test/sm.xml"])
        self.assertEqual(rb.clean_params, ["utm_source"])

    def test_robots_checks(self):
        self.assertEqual(self.res["robots_txt_valid"].status, WARN)  # нет Sitemap, есть устаревший Host
        self.assertEqual(self.res["robots_blocks_pages"].status, FAIL)  # /crm-test/ закрыт ($) и /Catalog/ в карте
        self.assertEqual(self.res["robots_vs_sitemap"].status, FAIL)
        self.assertIn("https://shop.test/Catalog/undefined/", self.urls("robots_vs_sitemap", ok=False))
        self.assertEqual(self.res["ai_bots_policy"].status, PASS)  # закрыт только GPTBot (обучение), не «ответный» бот
        self.assertEqual(self.res["yandex_clean_param"].status, WARN)  # ссылка с utm без Clean-param
        self.assertEqual(self.res["llms_txt"].status, SKIPPED)

    # ---- URL и ссылки
    def test_url_hygiene_and_anchors_and_mixed(self):
        self.assertEqual(self.res["url_hygiene"].status, FAIL)
        msg = next(i.msg for i in self.res["url_hygiene"].items if i.url == "https://shop.test/bread/")
        for word in ("заглавные", "метки utm", "http", "двойной слеш"):
            self.assertIn(word, msg)
        self.assertFalse(next(i for i in self.res["link_anchor_quality"].items if i.url == "https://shop.test/bread/").ok)
        self.assertEqual(self.res["mixed_content"].status, FAIL)
        self.assertIn("https://shop.test/bread/", self.urls("mixed_content", ok=False))

    def test_sitemap_hygiene_and_test_urls(self):
        r = self.res["sitemap_url_hygiene"]
        self.assertEqual(r.status, FAIL)
        for word in ("повторы URL", "не https", "чужой хост", "мусор"):
            self.assertIn(word, r.summary)
        t = self.res["test_urls_exposed"]
        self.assertEqual(t.status, FAIL)
        self.assertIn("https://shop.test/crm-test/", self.urls("test_urls_exposed"))
        self.assertNotIn("https://shop.test/milk/", self.urls("test_urls_exposed"))

    def test_crawl_depth_skipped_without_crawl(self):
        self.assertEqual(self.res["crawl_depth"].status, SKIPPED)

    # ---- сервер и заголовки
    def test_cache_control_parser(self):
        self.assertEqual(parse_cache_control("public, max-age=3600"), [])
        bad = parse_cache_control("public, max-age max-age=36000")
        self.assertTrue(any("слиты" in x for x in bad), bad)
        self.assertTrue(any("повтор" in x for x in parse_cache_control("public, public")))
        self.assertTrue(any("no-store" in x for x in parse_cache_control("no-store, max-age=60")))
        self.assertTrue(any("нечисл" in x for x in parse_cache_control("max-age=abc")))

    def test_header_checks(self):
        cc = {i.url: i for i in self.res["cache_headers"].items}
        self.assertTrue(cc["https://shop.test/milk/"].ok)
        self.assertFalse(cc["https://shop.test/bread/"].ok)
        self.assertEqual(self.res["compression"].status, FAIL)  # у b без Content-Encoding и страница больше 1,5 КБ
        self.assertTrue(next(i for i in self.res["compression"].items if i.url == "https://shop.test/milk/").ok)
        self.assertEqual(self.res["security_headers"].status, WARN)  # у большинства страниц заголовков нет
        v = self.res["vary_user_agent"]
        self.assertEqual(v.status, WARN)  # plain и bot различаются, Vary: User-Agent нет
        self.assertIn(self.res["bot_head_parity"].status, (WARN, FAIL))

    def test_html_weight_blocking_images_ok_on_small_pages(self):
        for cid in ("html_weight", "render_blocking", "images_perf"):
            self.assertIn(self.res[cid].status, (PASS, NA, WARN), cid)

    def test_tls_and_conditional_from_live_rows(self):
        st = self.store
        rows = [("tls", "https://shop.test/", 200, {"host": "shop.test", "ok": True, "days_left": 90, "issuer": "LE"}),
                ("tls", "https://www.shop.test/", None, {"host": "www.shop.test", "ok": False, "error": "самоподписанный сертификат",
                                                          "error_kind": "verify", "subject": "localhost"}),
                ("conditional", "https://shop.test/milk/", 200, {"validator": "ETag"})]
        for kind, url, status, meta in rows:
            st.db.execute("INSERT INTO live_urls(run_id,network,kind,url,status,final_url,location,hops,meta,fetched_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
                          ("r1", "shop", kind, url, status, url, None, 1, json.dumps(meta), now_iso()))
        st.db.commit()
        res = run_checks(st, self.settings, "r1", "shop", only={"tls_cert", "conditional_304"})
        self.assertEqual(res["tls_cert"].status, FAIL)
        self.assertIn("самоподписанный", " ".join(i.msg for i in res["tls_cert"].items))
        self.assertEqual(res["conditional_304"].status, WARN)  # 200 вместо 304 и страницы без валидаторов
        st.db.execute("DELETE FROM live_urls WHERE run_id='r1'")
        st.db.commit()

    def test_tls_info_closed_port_does_not_crash(self):
        r = tls_info("127.0.0.1", 1, 1.0)
        self.assertFalse(r["ok"])
        self.assertEqual(r["error_kind"], "connect")

    # ---- приоритет, роль, часы
    def test_priority_normalization_and_plan(self):
        self.assertEqual(plan.normalize_priority("{{USERNAME}}", "redirect_www"), "Критично")
        self.assertEqual(plan.normalize_priority("критично", "x"), "Критично")
        self.assertEqual(plan.normalize_priority("", "x"), "Средний")
        self.assertEqual(plan.normalize_priority("Низкий", "x"), "Низкий")
        rows = plan.enrich([{"block": "Canonical", "title": "t", "priority": "Критично", "check_id": "canonical_self"},
                            {"block": "Редиректы", "title": "u", "priority": "{{USERNAME}}", "check_id": "redirect_www"},
                            {"block": "Редиректы", "title": "ручной", "priority": "Высокий", "check_id": None}])
        self.assertEqual((rows[0]["role"], rows[0]["hours"], rows[0]["wave"]), ("DEV", 3.0, 1))
        self.assertEqual(rows[1]["priority"], "Критично")
        self.assertEqual(rows[1]["priority_raw"], "{{USERNAME}}")
        self.assertEqual(rows[2]["role"], "DEVOPS")  # по блоку
        self.assertEqual(plan.hours_for(4, "fail"), 4)
        self.assertEqual(plan.hours_for(4, "warn"), 2)
        self.assertEqual(plan.hours_for(4, "pass"), 0)

    def test_plan_file_overrides(self):
        f = Path(self.tmp.name) / "plan.csv"
        f.write_text("check_id;role;hours;result\ncanonical_self;SEO;10,5;Свой результат\n", encoding="utf-8")
        os.environ["PLAN_FILE"] = str(f)
        try:
            p = plan.plan_for("canonical_self")
            self.assertEqual((p["role"], p["hours"], p["result"]), ("SEO", 10.5, "Свой результат"))
            self.assertEqual(plan.plan_for("canonical_exact")["role"], "DEV")
        finally:
            os.environ.pop("PLAN_FILE", None)

    def test_tasks_and_totals(self):
        rows = plan.enrich(with_extras([]))
        results = load_results(self.store, "r1")
        tasks = plan.build_tasks(rows, results)
        self.assertTrue(tasks)
        self.assertEqual([t["wave"] for t in tasks], sorted(t["wave"] for t in tasks))
        tot = plan.totals_by_role(tasks)
        self.assertAlmostEqual(sum(a["hours"] for a in tot), sum(t["hours"] for t in tasks), places=2)

    # ---- отчёты
    def test_extras_are_added_to_rows_xlsx_and_html(self):
        import openpyxl
        rows = with_extras([{"row": 2, "block": "Canonical", "title": "Присутствует на всех индексируемых страницах", "criterion": "",
                             "priority": "{{USERNAME}}", "check_id": "canonical_present"}])
        self.assertEqual(len(rows) - 1, len(EXTRA))
        self.assertTrue(any(r.get("extra") for r in rows))
        src = Path(self.tmp.name) / "list.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Блок", "Пункт проверки", "Критерий", "Приоритет", "Проверяющий", "Комментарий"])
        ws.append(["Canonical", "Присутствует на всех индексируемых страницах", "один тег", "{{USERNAME}}", "a@b", "мой коммент"])
        wb.save(src)
        results = load_results(self.store, "r1")
        dst = Path(self.tmp.name) / "out.xlsx"
        checklist.write_xlsx(src, dst, results, "r1")
        out = openpyxl.load_workbook(dst)
        ws2 = out.worksheets[0]
        self.assertEqual(ws2.cell(2, 6).value, "мой коммент")  # исходные ячейки не тронуты
        self.assertEqual(ws2.cell(2, 4).value, "{{USERNAME}}")
        self.assertEqual(ws2.max_row, 2 + len(EXTRA))  # дополнительные проверки дописаны ниже
        self.assertIn("План работ", out.sheetnames)
        header = [c.value for c in ws2[1]]
        for h in ("Приоритет (норм.)", "Роль", "Часы (оценка)", "Результат работ", "Волна"):
            self.assertIn(h, header)
        sheet = out["План работ"]
        self.assertEqual([c.value for c in sheet[1]][:4], ["Волна", "Приоритет", "Блок", "Задача"])
        self.assertTrue(any(c.value == "Итого по ролям" for row in sheet.iter_rows() for c in row))
        html = report.build_html("r1", rows, results, {})
        self.assertIn("План работ", html)
        self.assertIn("Роль · часы", html)
        self.assertIn("Robots.txt и Яндекс", html)


class ExtraEndToEndTest(unittest.TestCase):
    """collect_html.py → run_checks.py --live на мок-сайте: новые проверки работают на живых данных, отчёты содержат план работ."""

    def test_live_and_reports(self):
        import io
        from contextlib import redirect_stdout
        from tests import mock_site
        from techaudit.collect import main as collect_main
        from techaudit.checks.cli import main as checks_main
        srv, base = mock_site.start()
        tmp = tempfile.TemporaryDirectory()
        t = Path(tmp.name)
        (t / "registry").mkdir()
        (t / "registry" / "mock.csv").write_text("type,url,note\nmain,/,\ncatalog,/catalog/,\nssr,/ssr/,\n", encoding="utf-8")
        saved = dict(os.environ)
        os.environ.update({"NETWORKS": "mock", "NET_MOCK_URL": base, "NET_MOCK_TOKEN": mock_site.TOKEN,
                           "DATA_DIR": str(t / "data"), "REGISTRY_DIR": str(t / "registry"), "UA_PROFILE": "googlebot"})
        try:
            with redirect_stdout(io.StringIO()):
                self.assertEqual(collect_main(["--variants", "plain,bot", "--probes", "all", "--no-diff"]), 0)
                self.assertEqual(checks_main(["--live", "--checklist", str(t / "none.xlsx")]), 0)
            st = Store(t / "data")
            rid = st.db.execute("SELECT run_id FROM runs").fetchone()[0]
            res = load_results(st, rid)["mock"]
            self.assertEqual(set(res), set(REGISTRY))
            for cid, r in res.items():
                self.assertFalse(r["summary"].startswith("ошибка проверки"), (cid, r["summary"]))
            self.assertEqual(res["llms_txt"]["status"], WARN)  # мок отдаёт 404 на /llms.txt
            self.assertEqual(res["tls_cert"]["status"], NA)  # адрес по IP: сертификат не проверяется
            self.assertEqual(res["robots_txt_valid"]["status"], WARN)  # Sitemap по http
            self.assertIn(res["robots_blocks_pages"]["status"], (PASS,))
            html = next((t / "data" / "reports").glob("*-checks.html")).read_text(encoding="utf-8")
            self.assertIn("План работ", html)
            self.assertIn("Индексация и дубли", html)
            st.close()
        finally:
            srv.shutdown(); srv.server_close()
            os.environ.clear(); os.environ.update(saved)
            tmp.cleanup()


if __name__ == "__main__":
    unittest.main()
