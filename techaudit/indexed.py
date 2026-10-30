"""Страницы, которые поисковики держат в индексе: выгрузки из Яндекс Вебмастера и Google Search Console.

Источники (поле source): ywm — Яндекс Вебмастер, gsc — Google Search Console. Выгрузка (index_imports) — один снимок списка URL сети:

  * через API Вебмастера: «Страницы в поиске» (search-urls/in-search/samples). Токен — YANDEX_WEBMASTER_TOKEN («Доступ → Внешние сервисы»).
    Ограничение API Яндекса: не более 50 000 URL на сайт за выгрузку; для больших сайтов дополняйте файлом из интерфейса Вебмастера.
  * через API Search Console: у Google нет метода «список проиндексированных страниц», поэтому берём страницы, которые получали показы
    в поиске за последние N дней (Search Analytics, измерение page) — это нижняя граница «страниц в индексе Google». Сервисные аккаунты и
    ресурсы — как для остальных проверок GSC («Доступ → Сети»). Полный список «Проиндексировано» из отчёта «Индексирование страниц»
    можно загрузить файлом (CSV/XLSX), но интерфейс Google отдаёт в выгрузку не больше 1000 строк на причину.
  * из файла: CSV / XLSX / TXT с колонкой адресов (заголовок URL, Адрес, Страница, Page, Top pages …) или просто список URL по строке.

Сравнение с краулом — в orphans.py.
"""
from __future__ import annotations

import csv
import io
import re
import sqlite3
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import BinaryIO, Callable, Iterable, Iterator
from urllib.parse import parse_qsl, quote, unquote, urlencode, urlsplit

import httpx

from .netutil import host_of, same_site

SOURCES = {"ywm": "Яндекс Вебмастер", "gsc": "Google Search Console"}

SCHEMA = """
CREATE TABLE IF NOT EXISTS index_imports (
  id INTEGER PRIMARY KEY AUTOINCREMENT, network TEXT NOT NULL, source TEXT NOT NULL,
  origin TEXT NOT NULL DEFAULT 'file',          -- api | file
  label TEXT, source_file TEXT, taken_at TEXT NOT NULL, rows INTEGER DEFAULT 0, skipped_external INTEGER DEFAULT 0, note TEXT
);
CREATE INDEX IF NOT EXISTS ix_index_imports ON index_imports(network, source, taken_at);
CREATE TABLE IF NOT EXISTS index_pages (
  import_id INTEGER NOT NULL, url TEXT NOT NULL, norm TEXT NOT NULL, title TEXT, last_access TEXT,
  clicks INTEGER, impressions INTEGER
);
CREATE INDEX IF NOT EXISTS ix_index_pages_imp ON index_pages(import_id, norm);
CREATE TABLE IF NOT EXISTS orphan_probe (
  network TEXT NOT NULL, norm TEXT NOT NULL, url TEXT, status INTEGER, final_url TEXT, error TEXT, checked_at TEXT,
  PRIMARY KEY (network, norm)
);
"""


class IndexedError(Exception):
    pass


def ensure_schema(db: sqlite3.Connection) -> None:
    db.executescript(SCHEMA)


# ---------------------------------------------------------------- нормализация адресов

TRACKING = re.compile(r"^(utm_.*|yclid|gclid|fbclid|ysclid|_openstat|from|openstat|_ga|_gl|mc_cid|mc_eid)$", re.I)


def normalize(url: str) -> str:
    """Ключ для сравнения адресов из разных источников: без схемы, www, порта по умолчанию, фрагмента, меток и закрывающего слеша.

    http/https считаем одним адресом (сайт редиректит одно на другое — это отдельная проверка); регистр хоста неважен, регистр пути сохраняется;
    процентное кодирование снимается (кириллица в Вебмастере и в краулe приходит по-разному).
    """
    u = (url or "").strip()
    if not u:
        return ""
    if "://" not in u:
        u = "https://" + u.lstrip("/")
    try:
        s = urlsplit(u)
    except ValueError:
        return u.lower()
    host = (s.hostname or "").lower().rstrip(".")
    if host.startswith("www."):
        host = host[4:]
    port = s.port if s.port not in (None, 80, 443) else None
    path = unquote(s.path or "/")
    path = re.sub(r"/{2,}", "/", path)
    if len(path) > 1:
        path = path.rstrip("/")
    q = sorted((k, v) for k, v in parse_qsl(s.query, keep_blank_values=True) if not TRACKING.match(k))
    out = host + (f":{port}" if port else "") + path
    if q:
        out += "?" + urlencode(q, safe="/:,")
    return out


def section_of(norm: str) -> str:
    """Первый сегмент пути — для группировки («/catalog», «/promo»); корень — «/»."""
    path = norm.split("?", 1)[0]
    parts = path.split("/", 2)
    return "/" + parts[1] if len(parts) > 1 and parts[1] else "/"


# ---------------------------------------------------------------- запись выгрузки

def new_import(db: sqlite3.Connection, network: str, source: str, *, origin: str, label: str = "", source_file: str = "",
               note: str = "") -> int:
    if source not in SOURCES:
        raise IndexedError(f"неизвестный источник «{source}»")
    ensure_schema(db)
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    cur = db.execute("INSERT INTO index_imports(network, source, origin, label, source_file, taken_at, rows, note) VALUES (?,?,?,?,?,?,0,?)",
                     (network, source, origin, label, source_file, now, note))
    return cur.lastrowid


def save_pages(db: sqlite3.Connection, network: str, source: str, records: Iterable[dict], *, host: str | None, origin: str,
               label: str = "", source_file: str = "", note: str = "", batch: int = 5000) -> dict:
    """records: {url, title?, last_access?, clicks?, impressions?}. Чужие хосты (если задан host) отбрасываются, дубли по ключу схлопываются."""
    ensure_schema(db)
    imp = new_import(db, network, source, origin=origin, label=label, source_file=source_file, note=note)
    seen: set[str] = set()
    buf: list[tuple] = []
    skipped = 0
    db.execute("BEGIN")
    try:
        for r in records:
            url = (r.get("url") or "").strip()
            if not url or not re.match(r"^(https?://|[\w.-]+\.[a-z]{2,}/|/)", url, re.I):
                continue
            if url.startswith("/"):  # путь без хоста — достраиваем хостом сети
                if not host:
                    continue
                url = f"https://{host}{url}"
            if host and not same_site(host_of(url if "://" in url else "https://" + url), host):
                skipped += 1
                continue
            n = normalize(url)
            if not n or n in seen:
                continue
            seen.add(n)
            buf.append((imp, url, n, (r.get("title") or None), (r.get("last_access") or None),
                        _int(r.get("clicks")), _int(r.get("impressions"))))
            if len(buf) >= batch:
                db.executemany("INSERT INTO index_pages VALUES (?,?,?,?,?,?,?)", buf)
                buf.clear()
        if buf:
            db.executemany("INSERT INTO index_pages VALUES (?,?,?,?,?,?,?)", buf)
        if not seen:
            raise IndexedError("в выгрузке нет адресов сети" + (f" (хост {host}); возможно, выбрана не та сеть" if skipped else ""))
        db.execute("UPDATE index_imports SET rows=?, skipped_external=? WHERE id=?", (len(seen), skipped, imp))
        db.execute("COMMIT")
    except BaseException:
        db.execute("ROLLBACK")
        db.execute("DELETE FROM index_imports WHERE id=?", (imp,))
        raise
    return {"import_id": imp, "rows": len(seen), "skipped_external": skipped}


def _int(v):
    try:
        return int(float(str(v).replace(",", ".").replace(" ", ""))) if v not in (None, "") else None
    except ValueError:
        return None


def imports(db: sqlite3.Connection, network: str | None = None, source: str | None = None, limit: int = 100) -> list[dict]:
    ensure_schema(db)
    where, args = [], []
    if network:
        where.append("network=?")
        args.append(network)
    if source:
        where.append("source=?")
        args.append(source)
    q = "SELECT * FROM index_imports" + (" WHERE " + " AND ".join(where) if where else "") + " ORDER BY taken_at DESC, id DESC LIMIT ?"
    return [dict(r) for r in db.execute(q, [*args, limit])]


def latest(db: sqlite3.Connection, network: str, source: str) -> dict | None:
    r = imports(db, network, source, 1)
    return r[0] if r else None


def delete_import(db: sqlite3.Connection, import_id: int) -> None:
    db.execute("BEGIN")
    db.execute("DELETE FROM index_pages WHERE import_id=?", (import_id,))
    db.execute("DELETE FROM index_imports WHERE id=?", (import_id,))
    db.execute("COMMIT")


# ---------------------------------------------------------------- файлы

URL_HEAD = re.compile(r"^(url|urls|адрес|адрес страницы|url страницы|страница|страницы|page|pages|address|landing page|"
                      r"самые популярные страницы|top pages|ссылка)$", re.I)
TITLE_HEAD = re.compile(r"^(title|заголовок|название|title 1)$", re.I)
DATE_HEAD = re.compile(r"^(last_access|last access|дата|дата обхода|последний обход|последнее посещение|дата посещения|"
                       r"last crawled|last crawl|последнее сканирование)$", re.I)
CLICKS_HEAD = re.compile(r"^(clicks|клики)$", re.I)
IMPR_HEAD = re.compile(r"^(impressions|показы)$", re.I)


def _pick(header: list[str], rx: re.Pattern) -> int | None:
    for i, h in enumerate(header):
        if rx.match((h or "").strip().lstrip("﻿")):
            return i
    return None


def _records_from_rows(rows: Iterator[list]) -> Iterator[dict]:
    first = next(rows, None)
    if first is None:
        raise IndexedError("файл пуст")
    cells = [str(c).strip() if c is not None else "" for c in first]
    ui = _pick(cells, URL_HEAD)
    if ui is None:
        # заголовка нет: адрес — первая ячейка, начинающаяся с http или с «/»
        ui = next((i for i, c in enumerate(cells) if re.match(r"^(https?://|/)", c, re.I)), None)
        if ui is None:
            raise IndexedError("не нашёл колонку с адресами: нужен заголовок URL / Адрес / Страница / Page или список адресов в первой колонке. "
                               "Найдены: " + ", ".join(cells[:8]))
        rest, header = [first], None
    else:
        rest, header = [], cells
    ti = _pick(header, TITLE_HEAD) if header else None
    di = _pick(header, DATE_HEAD) if header else None
    ci = _pick(header, CLICKS_HEAD) if header else None
    ii = _pick(header, IMPR_HEAD) if header else None

    def get(r, i):
        return str(r[i]).strip() if i is not None and i < len(r) and r[i] is not None else None

    for r in [*rest, *rows]:
        u = get(r, ui)
        if u:
            yield {"url": u, "title": get(r, ti), "last_access": get(r, di), "clicks": get(r, ci), "impressions": get(r, ii)}


def read_file(source: BinaryIO | Path | str, filename: str = "") -> Iterator[dict]:
    """Записи {url, title, last_access, clicks, impressions} из CSV / TXT / XLSX."""
    close = False
    if isinstance(source, (str, Path)):
        filename = filename or Path(source).name
        source, close = open(source, "rb"), True
    try:
        name = filename.lower()
        if name.endswith(".xls") and not name.endswith(".xlsx"):
            raise IndexedError("формат .xls (старый Excel) не читается: пересохраните как .xlsx или .csv")
        if name.endswith((".xlsx", ".xlsm")):
            try:
                from openpyxl import load_workbook
            except ImportError as e:  # pragma: no cover
                raise IndexedError("для .xlsx нужен пакет openpyxl (pip install openpyxl)") from e
            data = source.read()
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            yield from _records_from_rows(([c for c in row] for row in ws.iter_rows(values_only=True)))
            return
        raw = source.read()
        text = raw.decode("utf-8-sig", errors="replace") if isinstance(raw, (bytes, bytearray)) else str(raw)
        if "\x00" in text[:2000]:
            try:
                text = raw.decode("utf-16")
            except UnicodeError:
                pass
        lines = text.splitlines()
        head = next((ln for ln in lines if ln.strip()), "")
        if not head:
            raise IndexedError("файл пуст")
        delim = max((";", "\t", ","), key=head.count) if any(d in head for d in (";", "\t", ",")) else ";"
        if name.endswith(".txt") or delim not in head:
            yield from _records_from_rows(([ln.strip()] for ln in lines if ln.strip()))
            return
        yield from _records_from_rows((r for r in csv.reader(io.StringIO(text), delimiter=delim) if r and any(c.strip() for c in r)))
    finally:
        if close:
            source.close()


def import_file(db: sqlite3.Connection, network: str, source: str, file, *, filename: str = "", host: str | None = None,
                label: str = "") -> dict:
    name = filename or (Path(file).name if isinstance(file, (str, Path)) else "")
    return save_pages(db, network, source, read_file(file, name), host=host, origin="file", label=label, source_file=name)


# ---------------------------------------------------------------- Яндекс Вебмастер (API v4)

class Ywm:
    API = "https://api.webmaster.yandex.net/v4"
    PAGE = 100          # максимум метода
    MAX_URLS = 50000    # ограничение метода

    def __init__(self, token: str, *, timeout: float = 40, client: httpx.Client | None = None, pause: float = 0.15):
        if not (token or "").strip():
            raise IndexedError("не задан токен Яндекс Вебмастера: «Доступ → Внешние сервисы → Токен Яндекс Вебмастера» "
                               "(OAuth-токен с правом webmaster:hostinfo)")
        self._h = {"Authorization": f"OAuth {token.strip()}"}
        self._c = client or httpx.Client(timeout=timeout)
        self.pause = pause
        self.requests = 0
        self._uid: int | None = None

    def _get(self, path: str, params: dict | None = None) -> dict:
        last = ""
        for attempt in range(4):
            self.requests += 1
            try:
                r = self._c.get(self.API + path, headers=self._h, params=params)
            except httpx.HTTPError as e:
                last = type(e).__name__
                time.sleep(1.5 * (attempt + 1))
                continue
            if r.status_code == 200:
                return r.json()
            if r.status_code in (429, 500, 502, 503, 504):
                last = f"HTTP {r.status_code}"
                time.sleep(2.0 * (attempt + 1))
                continue
            msg = ""
            try:
                j = r.json()
                msg = f"{j.get('error_code', '')} {j.get('error_message', '')}".strip()
            except ValueError:
                msg = r.text[:150]
            if r.status_code in (401, 403):
                raise IndexedError(f"Вебмастер отклонил токен ({r.status_code} {msg}): проверьте токен и права доступа к сайту")
            raise IndexedError(f"Вебмастер ответил {r.status_code} {msg}")
        raise IndexedError(f"Вебмастер недоступен ({last})")

    def user_id(self) -> int:
        if self._uid is None:
            self._uid = int(self._get("/user")["user_id"])
        return self._uid

    def hosts(self) -> list[dict]:
        return self._get(f"/user/{self.user_id()}/hosts").get("hosts", [])

    def find_host(self, base_url: str) -> str:
        """host_id сайта в аккаунте: ищем по имени хоста (www не важен); при выборе предпочитаем https и подтверждённые права."""
        want = host_of(base_url)
        want = want[4:] if want.startswith("www.") else want
        cands = []
        for h in self.hosts():
            hid = h.get("host_id") or ""
            parts = hid.split(":")
            name = (parts[1] if len(parts) >= 2 else "").lower()
            name = name[4:] if name.startswith("www.") else name
            if name == want:
                cands.append((h.get("verified") is not False, parts[0] == "https", hid))
        if not cands:
            raise IndexedError(f"сайт {want} не найден среди ваших хостов в Вебмастере (или у токена нет доступа). "
                               "Можно указать host_id вручную, например https:5ka.ru:443")
        return sorted(cands, reverse=True)[0][2]

    def in_search(self, host_id: str, progress: Callable[[int, int], None] | None = None) -> Iterator[dict]:
        uid = self.user_id()
        path = f"/user/{uid}/hosts/{quote(host_id, safe='')}/search-urls/in-search/samples"
        offset, total = 0, None
        while offset < self.MAX_URLS:
            j = self._get(path, {"offset": offset, "limit": self.PAGE})
            total = j.get("count", total)
            samples = j.get("samples") or []
            if not samples:
                break
            for s in samples:
                yield {"url": s.get("url"), "title": s.get("title"), "last_access": s.get("last_access")}
            offset += len(samples)
            if progress:
                progress(offset, min(total or offset, self.MAX_URLS))
            if total is not None and offset >= total:
                break
            time.sleep(self.pause)
        self.total_reported = total


def fetch_ywm(db: sqlite3.Connection, network: str, base_url: str, token: str, *, host_id: str = "", label: str = "",
              progress: Callable[[int, int], None] | None = None, client: httpx.Client | None = None) -> dict:
    api = Ywm(token, client=client)
    hid = host_id.strip() or api.find_host(base_url)
    recs = list(api.in_search(hid, progress))
    total = getattr(api, "total_reported", None)
    note = f"host_id {hid}"
    if total and total > len(recs):
        note += f"; в Вебмастере {total} страниц, API отдаёт не больше {Ywm.MAX_URLS}"
    res = save_pages(db, network, "ywm", recs, host=host_of(base_url), origin="api", label=label, note=note)
    res.update(host_id=hid, total_reported=total, truncated=bool(total and total > len(recs)), requests=api.requests)
    return res


# ---------------------------------------------------------------- Google Search Console (Search Analytics)

def fetch_gsc(db: sqlite3.Connection, network: str, base_url: str, *, days: int = 90, label: str = "",
              progress: Callable[[int, int], None] | None = None, max_rows: int = 200000) -> dict:
    """Страницы с показами в Google за последние days дней. Перебирает ресурсы и аккаунты сети, пока один не ответит."""
    from .external import ExternalError, Gsc, Skipped
    g = Gsc(network, base_url)
    try:
        g._need_accounts()
    except Skipped as e:
        raise IndexedError(str(e)) from e
    end = date.today() - timedelta(days=2)  # свежие данные Google отдаёт с задержкой
    start = end - timedelta(days=max(1, days))
    errors: list[str] = []
    for account in g.accounts:
        for prop in g.properties:
            try:
                token = g._token(account)
            except (ExternalError, Skipped) as e:
                errors.append(str(e))
                break
            recs, err = [], None
            row = 0
            url = f"https://www.googleapis.com/webmasters/v3/sites/{quote(prop, safe='')}/searchAnalytics/query"
            while row < max_rows:
                try:
                    r = httpx.post(url, headers={"Authorization": f"Bearer {token}"}, timeout=90,
                                   json={"startDate": start.isoformat(), "endDate": end.isoformat(), "dimensions": ["page"],
                                         "type": "web", "rowLimit": 25000, "startRow": row})
                except httpx.HTTPError as e:
                    err = f"{type(e).__name__}"
                    break
                if r.status_code != 200:
                    err = f"{prop}: {r.status_code}"
                    break
                rows = r.json().get("rows") or []
                recs += [{"url": x["keys"][0], "clicks": x.get("clicks"), "impressions": x.get("impressions")} for x in rows]
                row += len(rows)
                if progress:
                    progress(row, row)
                if len(rows) < 25000:
                    break
            if err:
                errors.append(err)
                continue
            if recs:
                note = f"ресурс {prop}; страницы с показами за {days} дн. ({start} — {end})"
                res = save_pages(db, network, "gsc", recs, host=host_of(base_url), origin="api", label=label, note=note)
                res.update(property=prop, days=days)
                return res
            errors.append(f"{prop}: нет данных")
    raise IndexedError("Search Console не вернул страницы: " + ("; ".join(dict.fromkeys(errors)) or "нет ресурсов") +
                       ". Проверьте, что сервисный аккаунт добавлен в ресурс (права «Ограниченный» достаточно) и ресурс указан верно.")
