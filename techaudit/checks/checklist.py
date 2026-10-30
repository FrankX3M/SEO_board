"""Чек-лист Excel: чтение пунктов, сопоставление с проверками, запись результатов в копию файла."""
from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .core import FAIL, NA, PASS, SKIPPED, STATUS_LABELS, STATUS_ORDER, WARN, REGISTRY
from . import plan as planmod
from .extras import with_extras

FILLS = {PASS: "C6EFCE", FAIL: "FFC7CE", WARN: "FFEB9C", SKIPPED: "D9D9D9", NA: "EDEDED"}
COL_TITLE, COL_BLOCK = 2, 1


def find_checklist(root: Path, explicit: str | None = None) -> Path | None:
    cands = [explicit, os.environ.get("CHECKLIST_XLSX")]
    for c in cands:
        if c:
            p = Path(c)
            p = p if p.is_absolute() else root / p
            if p.exists():
                return p
    for d in (root, root.parent):
        for f in sorted(d.glob("*.xlsx")):
            if "тсх" in f.name.lower() or "чек" in f.name.lower() or "checklist" in f.name.lower():
                return f
    return None


def match_check(title: str) -> str | None:
    t = (title or "").lower()
    hit = [(cid, len(sub)) for cid, d in REGISTRY.items() for sub in d.titles if sub in t]
    return max(hit, key=lambda x: x[1])[0] if hit else None


def read_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    rows = []
    for r in range(2, ws.max_row + 1):
        title = ws.cell(r, COL_TITLE).value
        if not title:
            continue
        rows.append({"row": r, "block": ws.cell(r, COL_BLOCK).value, "title": str(title).strip(),
                     "criterion": ws.cell(r, 3).value, "priority": ws.cell(r, 4).value,
                     "check_id": match_check(str(title))})
    return rows


def write_plan_sheet(wb, rows: list[dict], results: dict, head_font) -> None:
    """Лист «План работ»: задачи по пунктам, не прошедшим проверку, с приоритетом, ролью, часами; итоги по ролям."""
    tasks = planmod.build_tasks(rows, results)
    sh = wb.create_sheet("План работ")
    sh.append(["Волна", "Приоритет", "Блок", "Задача", "Статус", "Сети", "Что найдено", "Роль", "Часы", "Результат работ"])
    for c in sh[1]:
        c.font = head_font
    for t in tasks:
        sh.append([planmod.WAVE_LABEL[t["wave"]], t["priority"], t["block"], t["title"], STATUS_LABELS[t["status"]], ", ".join(t["nets"]),
                   t["summary"], t["role_label"], t["hours"], t["result"]])
        sh.cell(sh.max_row, 5).fill = PatternFill("solid", fgColor=FILLS[t["status"]])
    last = sh.max_row
    sh.append([])
    sh.append(["Итого по ролям", "", "", "", "", "", "", "Роль", "Часы", "Задач"])
    for c in sh[sh.max_row]:
        c.font = head_font
    for a in planmod.totals_by_role(tasks):
        sh.append(["", "", "", "", "", "", "", a["role_label"], a["hours"], a["tasks"]])
    sh.append(["", "", "", "", "", "", "Всего", "", round(sum(t["hours"] for t in tasks), 2), len(tasks)])
    for c in sh[sh.max_row]:
        c.font = head_font
    sh.append([])
    sh.append(["Часы — оценка на исправление при статусе «Не пройдено»; при «Требует внимания» берётся "
               f"{planmod.warn_share():.0%} от оценки. Свои значения: PLAN_FILE в .env (CSV: check_id;role;hours;result)."])
    for col, w in zip("ABCDEFGHIJ", (10, 12, 26, 52, 18, 18, 70, 18, 8, 52)):
        sh.column_dimensions[col].width = w
    for row in sh.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    sh.freeze_panes = "A2"
    sh.auto_filter.ref = f"A1:J{last}"


def worst(statuses: list[str]) -> str:
    return min(statuses, key=lambda s: STATUS_ORDER.get(s, 9)) if statuses else NA


def write_xlsx(src: Path, dst: Path, results: dict[str, dict[str, dict]], run_id: str, notes: dict | None = None) -> dict:
    """Копирует чек-лист и добавляет колонки автопроверки. Исходные ячейки не меняются."""
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    nets = sorted(results)
    rows = planmod.enrich(with_extras(read_rows(src)))
    nxt = ws.max_row + 1
    for r in rows:  # дополнительные проверки дописываются ниже исходных строк, исходные ячейки не меняются
        if r["row"] is None:
            r["row"] = nxt
            nxt += 1
            for col, key in enumerate(("block", "title", "criterion", "priority"), 1):
                ws.cell(r["row"], col, r[key] if key != "priority" else r["priority"])
    base = ws.max_column + 1
    notes = notes or {}
    headers = ["Автопроверка: итог"] + [h for n in nets for h in (f"{n}: статус", f"{n}: результат")] + ["Источник данных", "Прогон"]
    if notes:
        headers.append("ИИ: черновик комментария (проверьте)")
    plan_col = base + len(headers)
    headers += ["Приоритет (норм.)", "Роль", "Часы (оценка)", "Результат работ", "Волна"]
    head_font = Font(bold=True)
    for i, h in enumerate(headers):
        c = ws.cell(1, base + i, h)
        c.font = head_font
        c.alignment = Alignment(wrap_text=True, vertical="top")
    unmatched = []
    for row in rows:
        r, cid = row["row"], row["check_id"]
        if not cid:
            unmatched.append(row["title"])
            ws.cell(r, base, STATUS_LABELS[SKIPPED]).fill = PatternFill("solid", fgColor=FILLS[SKIPPED])
            ws.cell(r, base + 1 + 2 * len(nets) , "нет автоматической проверки (ручной пункт)")
            continue
        sts = []
        for k, n in enumerate(nets):
            res = results[n].get(cid)
            if not res:
                continue
            sts.append(res["status"])
            c1 = ws.cell(r, base + 1 + 2 * k, STATUS_LABELS[res["status"]])
            c1.fill = PatternFill("solid", fgColor=FILLS[res["status"]])
            c2 = ws.cell(r, base + 2 + 2 * k, res["summary"])
            c2.alignment = Alignment(wrap_text=True, vertical="top")
        total = worst(sts)
        ws.cell(r, plan_col, row["priority"])
        ws.cell(r, plan_col + 1, row["role_label"])
        ws.cell(r, plan_col + 2, planmod.hours_for(row["hours"], total, len(nets)) or None)
        ws.cell(r, plan_col + 3, row["result"]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, plan_col + 4, planmod.WAVE_LABEL[row["wave"]])
        c0 = ws.cell(r, base, STATUS_LABELS[total])
        c0.fill = PatternFill("solid", fgColor=FILLS[total])
        c0.font = Font(bold=True)
        srcs = {results[n][cid]["source"] for n in nets if cid in results[n]}
        ws.cell(r, base + 1 + 2 * len(nets), ", ".join(sorted(srcs)))
        ws.cell(r, base + 2 + 2 * len(nets), run_id)
        if notes:
            txt = "\n".join(f"{n}: {notes[n]['analysis']['comments'][cid]}" for n in nets
                            if n in notes and cid in ((notes[n].get("analysis") or {}).get("comments") or {}))
            if txt:
                c = ws.cell(r, base + 3 + 2 * len(nets), txt)
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(len(headers)):
        ws.column_dimensions[get_column_letter(base + i)].width = (18 if i == 0 or "статус" in headers[i] else
                                                                   60 if ("результат" in headers[i].lower() or headers[i].startswith("ИИ")) else 16)
    ws.freeze_panes = "C2"

    # лист «Сводка» и «Детали»
    for name in ("Автопроверка - сводка", "Автопроверка - детали", "План работ", "ИИ - сводка"):
        if name in wb.sheetnames:
            del wb[name]
    sm = wb.create_sheet("Автопроверка - сводка")
    sm.append(["Сеть"] + [STATUS_LABELS[s] for s in (PASS, WARN, FAIL, SKIPPED, NA)] + ["Всего"])
    for n in nets:
        cnt = {s: sum(1 for x in results[n].values() if x["status"] == s) for s in FILLS}
        sm.append([n] + [cnt[s] for s in (PASS, WARN, FAIL, SKIPPED, NA)] + [sum(cnt.values())])
    sm.append([])
    sm.append([f"Запуск: {run_id}", f"Сформировано: {datetime.now():%Y-%m-%d %H:%M}"])
    if unmatched:
        sm.append([])
        sm.append(["Пункты без автопроверки:"])
        for t in unmatched:
            sm.append([t])
    for c in sm[1]:
        c.font = head_font
    sm.column_dimensions["A"].width = 28
    dt = wb.create_sheet("Автопроверка - детали")
    dt.append(["Блок", "Пункт", "Сеть", "Статус", "Тип страницы / пробы", "URL", "Результат"])
    for c in dt[1]:
        c.font = head_font
    for row in rows:
        cid = row["check_id"]
        if not cid:
            continue
        for n in nets:
            res = results[n].get(cid)
            if not res:
                continue
            bad = [i for i in res["items"] if not i["ok"]]
            for it in (bad or res["items"][:3]):
                dt.append([row["block"], row["title"], n, STATUS_LABELS["pass" if it["ok"] else res["status"] if res["status"] in (FAIL, WARN) else FAIL],
                           it.get("kind", ""), it["url"], it["msg"]])
    for col, w in zip("ABCDEFG", (22, 48, 12, 18, 18, 70, 70)):
        dt.column_dimensions[col].width = w
    dt.freeze_panes = "A2"
    if notes:
        ai = wb.create_sheet("ИИ - сводка")
        ai.append(["Сеть", "Раздел", "Текст"])
        for c in ai[1]:
            c.font = head_font
        for n in nets:
            a = (notes.get(n) or {}).get("analysis") or {}
            model = (notes.get(n) or {}).get("model", "")
            if a.get("summary"):
                ai.append([n, f"Сводка (ИИ, {model})", a["summary"]])
            for rc in a.get("root_causes") or []:
                ai.append([n, "Первопричина", f"{rc.get('title', '')} → {rc.get('action', '')} [{', '.join(rc.get('checks') or [])}]"])
            for t in a.get("tasks") or []:
                ai.append([n, f"Задача ({t.get('priority', '')})", f"{t.get('title', '')}\n{t.get('description', '')}"])
            for d in ((notes.get(n) or {}).get("descriptions") or {}).get("pages") or []:
                ai.append([n, "Текстовое описание", f"{d.get('url', '')}: есть={d.get('has_description')}, скрыто={d.get('hidden')}. {d.get('comment', '')}"])
        ai.append([])
        ai.append(["", "", "Вывод ИИ — черновик: статусы проверок он не меняет, перед использованием проверьте вручную."])
        ai.column_dimensions["A"].width, ai.column_dimensions["B"].width, ai.column_dimensions["C"].width = 14, 30, 120
        for row in ai.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    write_plan_sheet(wb, rows, results, head_font)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return {"unmatched": unmatched}
